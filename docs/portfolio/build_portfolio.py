"""포트폴리오 PPTX 생성 (python-pptx).

수정은 만들어진 .pptx를 PowerPoint에서 직접 하면 된다. 이 스크립트는 처음 한 번 뼈대를 만드는 용도.
실행: python build_portfolio.py   (python-pptx, Pillow 필요)
"""
import copy
import os

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
STUDENT = "정우영"
OUT = os.path.join(ROOT, "docs", "portfolio", f"01_블렌더기초_{STUDENT}.pptx")
IMG_CACHE = os.path.join(os.environ.get("TEMP", ROOT), "pp_img")
REPO_URL = "https://github.com/dnfltkagudwp-sys/Purrcelain-Panic"

FONT = "맑은 고딕"
NAVY = RGBColor(0x1B, 0x2A, 0x4A)      # 지배색: 청화백자 남색
COBALT = RGBColor(0x2F, 0x6F, 0xB5)
ORANGE = RGBColor(0xF2, 0x8C, 0x28)    # 포인트: 치즈냥 주황
TEXT = RGBColor(0x1F, 0x24, 0x30)
MUTED = RGBColor(0x5F, 0x67, 0x76)
CARD = RGBColor(0xEE, 0xF3, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xDD, 0xE8)

THEME_COLORS = {"dk1": "1F2430", "lt1": "FFFFFF", "dk2": "1B2A4A", "lt2": "EEF3F9",
                "accent1": "F28C28", "accent2": "2F6FB5", "accent3": "7FA7D9", "accent4": "8B5E3C",
                "accent5": "4CAF7A", "accent6": "D9534F", "hlink": "2F6FB5", "folHlink": "6D5BA8"}

SW, SH = 13.333, 7.5
M = 0.5  # 여백

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
LAYOUT_TITLE_ONLY = prs.slide_layouts[5]
LAYOUT_BLANK = prs.slide_layouts[6]

# ------------------------------------------------------------------ 테마: 글꼴·색
theme_part = prs.slide_master.part.part_related_by(RT.THEME)
theme = etree.fromstring(theme_part.blob)
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
for tag in ("majorFont", "minorFont"):
    f = theme.find(f".//{{{A}}}{tag}")
    f.find(f"{{{A}}}latin").set("typeface", FONT)
    f.find(f"{{{A}}}ea").set("typeface", FONT)
clr = theme.find(f".//{{{A}}}clrScheme")
clr.set("name", "Purrcelain")
for name, hexv in THEME_COLORS.items():
    el = clr.find(f"{{{A}}}{name}")
    for c in list(el):
        el.remove(c)
    etree.SubElement(el, f"{{{A}}}srgbClr", val=hexv)
theme_part._blob = etree.tostring(theme, xml_declaration=True, encoding="UTF-8", standalone=True)


# ------------------------------------------------------------------ 이미지 준비 (용량 줄이기)
def img(rel, max_px=1900):
    """rel: 경로 또는 (경로, (왼, 위, 오, 아래) 비율 자르기)."""
    crop = None
    if isinstance(rel, tuple):
        rel, crop = rel
    src = os.path.join(ROOT, rel)
    os.makedirs(IMG_CACHE, exist_ok=True)
    tag = "" if crop is None else "_crop" + "_".join(str(int(c * 100)) for c in crop)
    dst = os.path.join(IMG_CACHE, rel.replace("/", "_").rsplit(".", 1)[0] + tag + ".jpg")
    if not os.path.exists(dst) or os.path.getmtime(dst) < os.path.getmtime(src):
        im = Image.open(src).convert("RGB")
        if crop:
            W, H = im.size
            im = im.crop((int(crop[0] * W), int(crop[1] * H), int(crop[2] * W), int(crop[3] * H)))
        im.thumbnail((max_px, max_px), Image.LANCZOS)
        im.save(dst, "JPEG", quality=88)
    return dst


# ------------------------------------------------------------------ 도우미
def set_font(run, size=None, color=None, bold=None, italic=None):
    f = run.font
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", FONT)
    if size:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic


def set_bullet(p, char="•", indent=0.22):
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", str(int(Inches(indent))))
    ppr.set("indent", str(-int(Inches(indent))))
    for t in ("a:buNone", "a:buChar", "a:buAutoNum"):
        e = ppr.find(qn(t))
        if e is not None:
            ppr.remove(e)
    bu = etree.SubElement(ppr, qn("a:buChar"))
    bu.set("char", char)


def text(slide, x, y, w, h, paras, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, margin=0.04, name=None, space_after=4, line=1.1):
    """paras: str | [str | dict(text|runs, size, color, bold, align, bullet, space_after)]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(para, str):
            para = {"text": para}
        p.alignment = para.get("align", align)
        p.space_after = Pt(para.get("space_after", space_after))
        p.line_spacing = para.get("line", line)
        if para.get("bullet"):
            set_bullet(p)
        runs = para.get("runs") or [(para.get("text", ""), {})]
        for rt, ro in runs:
            r = p.add_run()
            r.text = rt
            set_font(r, size=ro.get("size", para.get("size", size)), color=ro.get("color", para.get("color", color)),
                     bold=ro.get("bold", para.get("bold", bold)), italic=ro.get("italic"))
    return box


def rect(slide, x, y, w, h, fill=CARD, line=None, radius=0.06, name=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(0.5, radius / min(w, h))
    s.text_frame.text = ""
    return s


def badge(slide, x, y, label, d=0.42, fill=ORANGE, color=WHITE, size=14):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.name = "Badge"
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    set_font(r, size=size, color=color, bold=True)
    return s


def pic(slide, rel, x, y, w, h, align="center", valign="middle", border=None, radius=0.08, name=None):
    """박스 안에 비율 유지로 맞춰 넣고 모서리를 살짝 둥글게."""
    path = img(rel)
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px = x + {"left": 0, "center": (w - pw) / 2, "right": w - pw}[align]
    py = y + {"top": 0, "middle": (h - ph) / 2, "bottom": h - ph}[valign]
    p = slide.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))
    if name:
        p.name = name
    p.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    geom = p._element.spPr.find(qn("a:prstGeom"))
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for g in list(av):
        av.remove(g)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(min(50000, radius / min(pw, ph) * 100000))}")
    if border is not None:
        p.line.color.rgb = border
        p.line.width = Pt(3)
    return p, (px, py, pw, ph)


def table(slide, x, y, w, rows, col_w, size=12, header_fill=NAVY, row_h=0.36, first_col_bold=True, name=None):
    nr, nc = len(rows), len(rows[0])
    shape = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    if name:
        shape.name = name
    t = shape.table
    tot = sum(col_w)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(w * cw / tot)
    for r in range(nr):
        t.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = t.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (CARD if r % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            val = rows[r][c]
            runs = val if isinstance(val, list) else [(str(val), {})]
            for rt, ro in runs:
                run = p.add_run()
                run.text = rt
                set_font(run, size=ro.get("size", size), bold=ro.get("bold", r == 0 or (c == 0 and first_col_bold)),
                         color=ro.get("color", WHITE if r == 0 else TEXT))
    return t


page_no = [0]
STAGES = ["컨셉", "리소스", "정리 과정", "결과"]
stage_pages = {k: [] for k in range(4)}


def new_slide(title, tag=None, dark=False, stage=None):
    """tag: 평가 항목 배지 문구, stage: 0 컨셉 · 1 리소스 · 2 정리 과정 · 3 결과 (하단 진행 표시)."""
    page_no[0] += 1
    if stage is not None:
        stage_pages[stage].append(page_no[0])
    s = prs.slides.add_slide(LAYOUT_TITLE_ONLY)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY if dark else WHITE
    t = s.shapes.title
    tx = M
    if tag:
        tag_w = 0.36 + sum(0.25 if ord(ch) > 0x3000 else 0.13 for ch in tag)
        b = rect(s, M, 0.42, tag_w, 0.42, fill=ORANGE, radius=0.21, name="Tag")
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = tag
        set_font(r, size=13, color=WHITE, bold=True)
        tx = M + tag_w + 0.2
    t.left, t.top, t.width, t.height = Inches(tx), Inches(0.3), Inches(SW - tx - M), Inches(0.66)
    tf = t.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    set_font(r, size=30, color=WHITE if dark else NAVY, bold=True)
    if stage is not None:
        runs = []
        for i, name in enumerate(STAGES):
            cur = (i == stage)
            runs.append((f"{i + 1} {name}", {"bold": cur, "color": ORANGE if cur else MUTED}))
            if i < 3:
                runs.append(("   ›   ", {"color": LINE}))
        text(s, M, SH - 0.45, 8, 0.32, [{"runs": runs}], size=11, margin=0, name="Progress")
    if not dark:
        text(s, SW - M - 1, SH - 0.45, 1, 0.32, str(page_no[0]), size=10, color=MUTED, align=PP_ALIGN.RIGHT,
             margin=0, name="PageNumber")
    return s


def caption(slide, x, y, w, s, size=11, align=PP_ALIGN.LEFT, color=MUTED):
    return text(slide, x, y, w, 0.3, s, size=size, color=color, align=align, margin=0)


def stat(slide, x, y, w, big, label, big_color=ORANGE, big_size=34, label_size=12, h=1.15, fill=CARD, label_color=TEXT):
    rect(slide, x, y, w, h, fill=fill, radius=0.12)
    text(slide, x + 0.2, y + 0.1, w - 0.4, 0.6, big, size=big_size, color=big_color, bold=True, margin=0)
    text(slide, x + 0.2, y + h - 0.45, w - 0.4, 0.38, label, size=label_size, color=label_color, margin=0)


def bullets(slide, x, y, w, h, items, size=14, head=None):
    paras = [{"text": head, "size": size + 2, "bold": True, "color": NAVY, "space_after": 8}] if head else []
    paras += [{"text": it, "bullet": True} for it in items]
    return text(slide, x, y, w, h, paras, size=size, margin=0, space_after=6)


VARIANT_IMG = ("Captures/process/P10_도자기_변형체.png", (0.02, 0.24, 0.98, 0.76))

# =================================================================== 표지
s = new_slide("Purrcelain Panic", dark=True)
s.shapes.title.top, s.shapes.title.height = Inches(1.7), Inches(1.1)
s.shapes.title.width = Inches(6.2)
s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(48)
text(s, M, 2.85, 6.0, 1.0, [
    {"text": "블렌더 기초 — AI 에셋 게임용 정리 포트폴리오", "size": 20, "color": WHITE, "bold": True},
    {"text": "VARCO 3D 리소스를 Blender와 Claude로 정리해 Unity용 FBX로 출력", "size": 14,
     "color": RGBColor(0xCA, 0xD6, 0xE8)},
], margin=0)
text(s, M, 4.6, 6.0, 1.6, [
    {"runs": [("훈련생   ", {"color": RGBColor(0xCA, 0xD6, 0xE8)}), (STUDENT, {"color": WHITE, "bold": True})]},
    {"runs": [("과정   ", {"color": RGBColor(0xCA, 0xD6, 0xE8)}), ("2026 대전 KNDA 게임 개발 과정", {"color": WHITE})]},
    {"runs": [("저장소   ", {"color": RGBColor(0xCA, 0xD6, 0xE8)}), (REPO_URL, {"color": WHITE})]},
], size=14, margin=0, space_after=8)
pic(s, "Captures/after/layout_push8.png", 6.9, 1.2, 5.95, 5.1, radius=0.18, name="CoverImage")

# =================================================================== 한눈에 보기
overview = new_slide("한눈에 보기")
flow = [("게임 컨셉", "Captures/after/layout_idle.png", "치즈냥이 도자기를\n밀어 떨어뜨려 깨는 게임"),
        ("리소스", "Concept/cat_B.png", "VARCO 3D로\n고양이·식탁·도자기 생성"),
        ("정리 과정", "Captures/process/P2_열린엣지_병합_전후.png", "Blender + Claude로\n게임에 맞게 정리"),
        ("결과", "Captures/unity/Unity_Compare_BeforeAfter.png", "Unity용 FBX 5종\n정리 전·후 비교")]
fw, fgap = 2.75, 0.43
overview_cards = []
for i, (name, image, desc) in enumerate(flow):
    x = M + i * (fw + fgap)
    rect(overview, x, 1.3, fw, 5.35, fill=CARD, radius=0.16)
    badge(overview, x + 0.25, 1.5, str(i + 1), d=0.5, size=16)
    text(overview, x + 0.88, 1.5, fw - 1.0, 0.5, name, size=20, bold=True, color=NAVY, margin=0,
         anchor=MSO_ANCHOR.MIDDLE)
    pic(overview, image, x + 0.2, 2.25, fw - 0.4, 2.35, radius=0.1)
    text(overview, x + 0.25, 4.8, fw - 0.5, 1.0, desc, size=13, margin=0)
    overview_cards.append(x)
    if i < 3:
        a = overview.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + fw + 0.08), Inches(3.3), Inches(0.27),
                                      Inches(0.36))
        a.fill.solid()
        a.fill.fore_color.rgb = ORANGE
        a.line.fill.background()

# =================================================================== 1 컨셉
s = new_slide("게임 컨셉", tag="컨셉", stage=0)
pic(s, "Captures/after/layout_idle.png", M, 1.25, 7.3, 4.6, name="ConceptImage")
text(s, 8.2, 1.35, 4.6, 1.2, "식탁 뒤에 기대 선 치즈냥이 앞발로 도자기를 밀어 떨어뜨려 깨는 게임",
     size=18, bold=True, color=NAVY, margin=0)
bullets(s, 8.2, 2.65, 4.6, 1.4, ["도자기가 점점 많이 생성되고 전부 깨는 재미", "하체는 식탁에 가려져 걷기 대신 좌우 이동"],
        size=14)
for i, (name, desc) in enumerate([("고양이", "캐릭터 · 애니메이션 3종"), ("식탁", "배경 소품"), ("도자기", "파괴 대상 · 3종")]):
    y = 4.3 + i * 0.6
    badge(s, 8.2, y + 0.04, str(i + 1), d=0.4, size=13)
    text(s, 8.75, y, 4.0, 0.5, [{"runs": [(name + "  ", {"bold": True, "color": NAVY}), (desc, {"color": TEXT})]}],
         size=14, margin=0, anchor=MSO_ANCHOR.MIDDLE)

# =================================================================== 2 리소스
s = new_slide("VARCO 3D 리소스 생성 ① 컨셉 이미지", tag="AI 리소스", stage=1)
cols = [("고양이", "cat_A", "cat_B", 1), ("식탁", "table_A", "table_B", 0), ("도자기", "vase_A", "vase_B", 0)]
gx, gy, cell, gap = M, 1.55, 2.25, 0.25
for c, (label, a, b, sel) in enumerate(cols):
    x = gx + c * (cell + gap)
    text(s, x, 1.2, cell, 0.32, label, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER, margin=0)
    for r, key in enumerate((a, b)):
        y = gy + r * (cell + 0.2)
        chosen = (r == sel)
        pic(s, f"Concept/{key}.png", x, y, cell, cell, border=ORANGE if chosen else LINE, radius=0.12)
        if chosen:
            t = rect(s, x + 0.12, y + 0.12, 0.62, 0.3, fill=ORANGE, radius=0.15, name="Chosen")
            tf = t.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            rr = p.add_run()
            rr.text = "선택"
            set_font(rr, size=11, color=WHITE, bold=True)
rx = gx + 3 * cell + 2 * gap + 0.45
bullets(s, rx, 1.3, SW - rx - M, 4.5, [
    "에셋마다 2장씩 생성해 1장 선택",
    "공통 스타일: 카툰 · 장난감처럼 매끈한 형태",
    "고양이는 자동 리깅을 위해 T-pose",
    "식탁은 장식 없는 단순한 형태",
    "도자기는 흰색 + 파랑으로 잘 보이게",
], size=14, head="선택 기준")

s = new_slide("VARCO 3D 리소스 생성 ② 3D 메시와 자동 리깅", tag="AI 리소스", stage=1)
raw = [("Captures/raw/cat_rigged_front.png", "고양이", "50,000 tris · 자동 리깅"),
       ("Captures/raw/table_raw_front.png", "식탁", "20,000 tris"),
       ("Captures/raw/vase_raw_front.png", "도자기", "20,000 tris")]
cw = (SW - 2 * M - 2 * 0.35) / 3
for i, (path, name, l1) in enumerate(raw):
    x = M + i * (cw + 0.35)
    pic(s, path, x, 1.2, cw, 4.1, radius=0.12)
    text(s, x, 5.45, cw, 0.8, [{"text": name, "size": 16, "bold": True, "color": NAVY}, {"text": l1, "size": 13}],
         margin=0, space_after=2)
caption(s, M, 6.45, SW - 2 * M, "원본은 수정하지 않고 Source/에 보존 · 폴리곤을 넉넉히 생성하고 Blender에서 줄임", size=12)

# =================================================================== 3 정리 과정
s = new_slide("작업 방식: 요청 → bpy 스크립트 → Blender", tag="Claude 활용", stage=2)
steps = [("요청문 작성", "Claude에게 요청"), ("bpy 스크립트", "Claude가 작성"), ("Blender 실행", "원본은 백업 보존"),
         ("검사·렌더", "수치 + 비교 이미지"), ("FBX 출력", "다시 불러와 검증")]
bw, bgap = 2.15, 0.33
for i, (h1, h2) in enumerate(steps):
    x = M + i * (bw + bgap)
    rect(s, x, 1.5, bw, 1.6, fill=CARD, radius=0.14)
    badge(s, x + 0.18, 1.68, str(i + 1), d=0.42, size=13)
    text(s, x + 0.18, 2.2, bw - 0.3, 0.8, [{"text": h1, "size": 15, "bold": True, "color": NAVY},
                                           {"text": h2, "size": 12, "color": MUTED}], margin=0, space_after=2)
    if i < len(steps) - 1:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.06), Inches(2.15), Inches(0.21), Inches(0.3))
        a.fill.solid()
        a.fill.fore_color.rgb = ORANGE
        a.line.fill.background()
table(s, M, 3.6, SW - 2 * M, [
    ["스크립트", "하는 일"],
    ["00_import · 01_cat", "원본 백업, 고양이 정리와 애니메이션"],
    ["02_vase · 03b_vase_variants", "도자기 파편, 변형체 2종"],
    ["03_table_style", "식탁 정리, 머티리얼 톤 통일"],
    ["04_export", "최종 검수, FBX 출력과 재검증"],
], [3.6, 8.7], size=13, row_h=0.48)
caption(s, M, 6.15, SW - 2 * M, "Blender/scripts/ · run_pipeline.sh 한 번으로 원본부터 다시 생성", size=12)

s = new_slide("게임 규격 정리: 크기 · 원점 · 방향", tag="규격", stage=2)
pic(s, "Captures/process/P6_크기_원점.png", M, 1.2, 8.3, 4.15, name="ScaleImage")
caption(s, M, 5.45, 8.3, "같은 축척 정면 · 왼쪽 막대 1 m · 빨간 점 = 원점", size=12)
sx, sw_ = 9.15, SW - 9.15 - M
stat(s, sx, 1.2, sw_, "0.84 → 1.15 m", "고양이 높이", big_size=28)
stat(s, sx, 2.5, sw_, "1.80 × 0.60 × 0.62", "식탁 (m)", big_size=24)
stat(s, sx, 3.8, sw_, "발밑 · 바닥 중앙", "원점 위치", big_size=24)
stat(s, sx, 5.1, sw_, "Rot 0 · Scale 1", "Apply Transform", big_size=24)

s = new_slide("최적화: 폴리곤 줄이기", tag="최적화", stage=2)
pic(s, "Captures/process/P1_최적화_와이어프레임_전후.png", M, 1.15, 3.85, 5.75, align="left", name="WireImage")
sx = 4.85
for i, (big, lab) in enumerate([("−60%", "고양이 50k → 20k"), ("−75%", "식탁 20k → 5k"), ("−70%", "도자기 20k → 6k")]):
    stat(s, sx + i * 2.72, 1.2, 2.5, big, lab, big_size=36, label_size=13)
bullets(s, sx, 2.85, SW - sx - M, 3.5, [
    "갈라진 정점을 Merge by Distance로 병합",
    "Decimate로 실루엣이 유지되는 만큼만 줄임",
    "불필요한 조각(Icosphere) 삭제",
], size=15, head="방법")

s = new_slide("편집·검수 ① 열린 엣지(구멍) 수정", tag="편집·검수", stage=2)
pic(s, "Captures/process/P2_열린엣지_병합_전후.png", M, 1.15, 8.4, 5.6, align="left", name="OpenEdgeImage")
sx = 9.2
stat(s, sx, 1.2, SW - sx - M, "8,244 → 0", "고양이 열린 엣지", big_size=30)
stat(s, sx, 2.5, SW - sx - M, "4,636 → 0", "식탁 열린 엣지", big_size=30)
stat(s, sx, 3.8, SW - sx - M, "2,260 → 0", "도자기 열린 엣지", big_size=30)
bullets(s, sx, 5.2, SW - sx - M, 1.5, ["원본은 UV 이음매마다 갈라져 있었음", "노멀 검사: 뒤집힌 면 0"], size=13)

s = new_slide("편집·검수 ② 도자기 파츠 분리", tag="편집·검수", stage=2)
pic(s, "Captures/process/P5_도자기_파츠분리.png", M, 1.15, SW - 2 * M, 3.1, name="FractureImage")
cards = [("9조각", "작은 조각은 빼고 다시 잘라 9조각"),
         ("닫힌 메시", "구멍·뒤집힌 면 0 → Unity 물리 사용 가능"),
         ("같은 위치·재질", "원본과 겹쳐 확인 가능, 단면은 흰 도자기 색")]
cw = (SW - 2 * M - 2 * 0.3) / 3
for i, (h1, body) in enumerate(cards):
    x = M + i * (cw + 0.3)
    rect(s, x, 4.5, cw, 1.6, fill=CARD, radius=0.14)
    text(s, x + 0.25, 4.65, cw - 0.5, 1.35, [{"text": h1, "size": 22, "bold": True, "color": ORANGE},
                                            {"text": body, "size": 14}], margin=0, space_after=6)

s = new_slide("스타일 통일: 색감과 거칠기", tag="스타일", stage=2)
pic(s, "Captures/process/P4_머티리얼_톤통일_전후.png", M, 1.15, 7.9, 5.3, align="left", name="MaterialImage")
sx = 8.7
text(s, sx, 1.2, SW - sx - M, 0.9, "고양이는 가장 눈에 띄게, 식탁은 중간 밝기, 도자기는 밝게",
     size=16, bold=True, color=NAVY, margin=0)
table(s, sx, 2.3, SW - sx - M, [
    ["에셋", "채도", "밝기", "거칠기"],
    ["고양이", "↑", "↑", "0.65"],
    ["식탁", "↓", "중간", "0.65"],
    ["도자기", "↑", "↑", "0.35"],
], [1.3, 0.9, 1.1, 1.1], size=13, row_h=0.45)
bullets(s, sx, 4.35, SW - sx - M, 2.0, ["색 보정을 텍스처에 직접 적용 → Unity에서도 유지", "도자기만 광택(유약) 표현"], size=13)

s = new_slide("애니메이션: Idle_Lean · Push_R · Push_L", tag="애니메이션", stage=2)
for i, (path, name, info) in enumerate([
        ("Captures/process/P8_Idle_Lean_키포즈.png", "Idle_Lean", "30프레임 · 반복"),
        ("Captures/process/P7_Push_R_키포즈.png", "Push_R", "16프레임 · 1회"),
        ("Captures/process/P9_Push_L_키포즈.png", "Push_L", "Push_R 좌우 반전")]):
    y = 1.15 + i * 1.83
    pic(s, path, M, y, 7.0, 1.73, align="left")
    text(s, 7.75, y + 0.35, SW - 7.75 - M, 1.2, [{"text": name, "size": 20, "bold": True, "color": ORANGE},
                                                {"text": info, "size": 14}], margin=0, space_after=4)

# --- Claude 활용
s = new_slide("Claude 활용: 요청 기록", tag="Claude 활용", stage=2)
rect(s, M, 1.25, 4.0, 5.3, fill=NAVY, radius=0.16)
text(s, M + 0.35, 1.5, 3.3, 4.9, [
    {"text": "작업 브리프", "size": 20, "bold": True, "color": WHITE},
    {"text": "요청문 1~4를 한 번에 전달", "size": 13, "color": ORANGE, "space_after": 14},
    {"text": "1  고양이 정리와 애니메이션", "size": 14, "color": WHITE},
    {"text": "2  도자기 파편 제작", "size": 14, "color": WHITE},
    {"text": "3  식탁 정리와 스타일 통일", "size": 14, "color": WHITE},
    {"text": "4  최종 검수와 FBX 출력", "size": 14, "color": WHITE},
], margin=0, space_after=8)
text(s, 4.85, 1.25, 7.9, 0.4, "작업 중 결과를 보고 따로 요청한 편집", size=16, bold=True, color=NAVY, margin=0)
later = [("5", "식탁 넓히고 낮추기", "1.8 × 0.6 × 0.62 m, 팔 각도 맞춤"),
         ("6", "왼손 밀기 Push_L", "Push_R 좌우 반전"),
         ("7", "도자기 변형체 2종", "늘린 청자 · 누른 분홍")]
for i, (n, h1, res) in enumerate(later):
    y = 1.8 + i * 1.6
    rect(s, 4.85, y, 7.98, 1.4, fill=CARD, radius=0.14)
    badge(s, 5.1, y + 0.45, n, d=0.5, size=16)
    text(s, 5.85, y + 0.2, 6.8, 1.0, [{"text": f"요청문 {n} · {h1}", "size": 17, "bold": True, "color": NAVY},
                                      {"text": res, "size": 13}], margin=0, space_after=4, anchor=MSO_ANCHOR.MIDDLE)


def request_slide(n, title, quote, results, image):
    s = new_slide(f"요청문 {n} · {title}", tag="Claude 활용", stage=2)
    rect(s, M, 1.2, 6.2, 2.6, fill=NAVY, radius=0.16)
    text(s, M + 0.35, 1.35, 5.5, 2.3, [{"text": "요청문 원문", "size": 12, "bold": True, "color": ORANGE},
                                       {"text": quote, "size": 18, "color": WHITE}],
         margin=0, space_after=8, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, M, 4.1, 6.2, 2.6, results, size=15, head="결과")
    pic(s, image, 7.05, 1.2, SW - 7.05 - M, 5.5, name="RequestImage")


request_slide(5, "식탁 넓히고 낮추기",
              "“책상 좀 넓히고 높이 좀 낮출 수 있나? 이 게임은 도자기가 책상에 점점 많이 생성되고 그걸 다 깨는 재미인데 책상이 너무 작네. 고양이 팔이 안닿으려나”",
              ["식탁 1.24 × 0.87 × 0.72 → 1.80 × 0.60 × 0.62 m", "다리는 그대로 두고 가운데만 늘림",
               "낮아진 식탁에 맞춰 팔 각도 자동 조정"],
              "Captures/after/layout_push8.png")
request_slide(6, "왼손 밀기 Push_L",
              "“왼손 PUSH_1도 만들어줘. PUSH_R이랑 동일하게, 손만 바꿔서 만들어주면 돼.”",
              ["Push_R을 좌우 반전해 Push_L 생성", "길이·타이밍 동일 (16프레임)", "Unity 클립 3개"],
              "Captures/process/P9_Push_L_키포즈.png")
request_slide(7, "도자기 변형체 2종",
              "“그 도자기 같은 것만 나오면 너무 지루하니까, 색깔과 모양이 조금씩 다른 변형체를 2개만 추가해줘. 조금 늘리거나 누르는 식으로.”",
              ["늘린 청자(초록) · 누른 분홍 2종", "파편도 같이 변형해 그대로 깨짐", "FBX 2종 추가"],
              VARIANT_IMG)

s = new_slide("시행착오와 해결", tag="과정", stage=2)
table(s, M, 1.3, SW - 2 * M, [
    ["문제", "해결"],
    ["도자기 파편이 6조각만 생성됨", "위치 계산을 갱신한 뒤 다시 자름"],
    ["작은 파편을 합치자 메시가 깨짐", "합치지 않고 그 조각을 빼고 다시 자름"],
    ["Unity 클립 이름이 CAT_Rig|Idle_Lean", "출력 방식을 바꿔 Idle_Lean으로 정리"],
    ["원본 백업의 색까지 바뀜", "백업 머티리얼을 따로 복사"],
    ["낮춘 식탁에서 앞발이 떠 있음", "식탁 높이에 맞춰 팔 각도 자동 계산"],
], [6.0, 6.3], size=15, row_h=0.82)

# =================================================================== 4 결과
s = new_slide("Unity에서 정리 전·후 비교", tag="결과 검증", stage=3)
pic(s, "Captures/unity/Unity_Compare_BeforeAfter.png", M, 1.15, 8.9, 4.45, align="left", name="UnityImage")
caption(s, M, 5.7, 8.9, "Assets/Scenes/Compare_BeforeAfter.unity", size=12)
sx = 9.7
text(s, sx, 1.2, SW - sx - M, 5.4, [
    {"text": "정리 전", "size": 17, "bold": True, "color": RGBColor(0xC0, 0x39, 0x2B)},
    {"text": "원점이 중앙 → 바닥에 묻힘", "bullet": True},
    {"text": "실제 비율 아님", "bullet": True},
    {"text": "90,000 tris", "bullet": True, "space_after": 16},
    {"text": "정리 후", "size": 17, "bold": True, "color": RGBColor(0x2E, 0x7D, 0x4F)},
    {"text": "원점 바닥 · Rot 0 · Scale 1", "bullet": True},
    {"text": "실제 비율", "bullet": True},
    {"text": "31,000 tris", "bullet": True},
], size=14, margin=0, space_after=5)

s = new_slide("최종 결과물: Unity용 FBX 5종", tag="결과", stage=3)
table(s, M, 1.25, 7.6, [
    ["파일", "내용", "tris"],
    ["Cat_Player.fbx", "고양이 + Idle_Lean · Push_R · Push_L", "20,000"],
    ["Table_Main.fbx", "식탁", "5,000"],
    ["Vase_Set.fbx", "도자기 + 파편 9조각", "6,000"],
    ["Vase_Tall_Set.fbx", "늘린 청자 + 파편 9조각", "6,000"],
    ["Vase_Squat_Set.fbx", "누른 분홍 + 파편 9조각", "6,000"],
], [2.4, 4.0, 1.2], size=14, row_h=0.72)
caption(s, M, 5.7, 7.6, "도자기 tris는 원본 기준 (파편 묶음은 각 8,500 별도)", size=12)
pic(s, "Captures/after/layout_push8.png", 8.45, 1.25, SW - 8.45 - M, 2.6, valign="top", name="FinalImage")
pic(s, VARIANT_IMG, 8.45, 4.05, SW - 8.45 - M, 2.6, valign="top", name="VariantImage")

s = new_slide("정리와 다음 단계", dark=True, stage=3)
for i, (big, lab) in enumerate([("5종", "FBX (캐릭터 1 · 소품 4)"), ("90k → 31k", "폴리곤 (파편 제외)"),
                                 ("7건", "Claude 요청문"), ("0", "구멍 · 뒤집힌 면")]):
    x = M + i * 3.13
    stat(s, x, 1.4, 2.9, big, lab, big_size=32, label_size=12, h=1.3, fill=RGBColor(0x26, 0x3A, 0x63),
         label_color=RGBColor(0xDD, 0xE6, 0xF2))
text(s, M, 3.2, SW - 2 * M, 3.0, [
    {"text": "2차 Unity 평가에서 할 일", "size": 18, "bold": True, "color": ORANGE, "space_after": 10},
    {"text": "고양이 좌우 이동 + 양손 밀기 입력 연결", "bullet": True},
    {"text": "도자기 3종을 식탁 위에 점점 많이 생성", "bullet": True},
    {"text": "파편에 충돌체와 물리를 붙여 깨지는 연출", "bullet": True},
], size=15, color=WHITE, margin=0, space_after=8)
text(s, M, 6.55, SW - 2 * M, 0.35, REPO_URL, size=12, color=RGBColor(0xCA, 0xD6, 0xE8), margin=0)

# --- 한눈에 보기: 단계별 쪽 번호
for i, x in enumerate(overview_cards):
    pages = stage_pages[i]
    label = f"{pages[0]}쪽" if len(pages) == 1 else f"{pages[0]}–{pages[-1]}쪽"
    text(overview, x + 0.25, 5.95, 2.25, 0.4, label, size=12, bold=True, color=ORANGE, margin=0)

# =================================================================== 저장
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides", round(os.path.getsize(OUT) / 1e6, 1), "MB")
